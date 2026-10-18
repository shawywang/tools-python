import os
from typing import Set, Dict, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string


# pip install openpyxl
# C:\ProgramData\miniconda3\python.exe -m pip install --upgrade pip ; C:\ProgramData\miniconda3\python.exe -m pip install openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, "筛选配置.txt")

# 筛选目标名称 -> 原脚本里的参数字段
FILTER_MAP = {
    "包含则提取": "inc",
    "包含则剔除": "excl",
    "等于则剔除": "excl_eq",
}


def parse_config(path):
    """读取 txt 配置，按缩进判断层级，返回 (全局配置, 用户字典)。"""
    with open(path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    config: Dict[str, str] = {}
    users = {}  # 用户名 -> {"_sheets": [sheet, ...]}
    stack = []  # 每层记录 (缩进, 类型, 对象)，用于向上回溯

    def close_to(indent):
        while stack and stack[-1][0] >= indent:
            stack.pop()

    for line in lines:
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        indent = len(line) - len(line.lstrip(" \t"))
        close_to(indent)

        if indent == 0:
            # 顶层：全局键值 或 新用户
            if ":" in s:
                key, _, val = s.partition(":")
                config[key.strip()] = val.strip()
                continue
            users.setdefault(s, {"_sheets": []})
            stack.append((indent, "user", users[s]))
            continue

        parent = stack[-1]
        if parent[1] == "user":
            # 工作表
            sheet = {"name": s, "skip_rows": 0,
                     "inc": {}, "excl": {}, "excl_eq": {}}
            parent[2]["_sheets"].append(sheet)
            stack.append((indent, "sheet", sheet))
        elif parent[1] == "sheet":
            # 跳过行数 或 筛选目标
            if s in FILTER_MAP:
                stack.append((indent, "filter", parent[2][FILTER_MAP[s]]))
            elif s.startswith("跳过行数"):
                parent[2]["skip_rows"] = int(s.partition(":")[2].strip())
        elif parent[1] == "filter":
            # 具体配置项：列名: 关键字（列名带「列」字，去掉后即 Excel 列标）
            col, _, kw = s.partition(":")
            col = col.strip().rstrip("列")
            kw = kw.strip()
            if col and kw:
                parent[2].setdefault(col, set()).add(kw)

    return config, users


class Handle:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.wb: Optional[Workbook] = None

    def del_column(self, sh_name: str, skip_rows: int, inc: Dict[str, Set[str]], excl: Dict[str, Set[str]], excl_eq: Dict[str, Set[str]]):
        ws = self.wb[sh_name]
        rows = list(ws.rows)  # 读成list
        ws.delete_rows(1, ws.max_row)  # 清空sheet

        for k, v in inc.items():
            col_inc_n = column_index_from_string(k) - 1
            row_n = skip_rows
            while row_n < len(rows):
                t = rows[row_n][col_inc_n].value
                if not t:
                    row_n += 1
                    continue
                if not any(s in t for s in v):
                    del rows[row_n]
                else:
                    row_n += 1

        for k, v in excl.items():
            col_excel_n = column_index_from_string(k) - 1
            row_n = skip_rows
            while row_n < len(rows):
                t = rows[row_n][col_excel_n].value
                if not t:
                    row_n += 1
                    continue
                if any(s in t for s in v):
                    del rows[row_n]
                else:
                    row_n += 1

        for k, v in excl_eq.items():
            col_excel_n = column_index_from_string(k) - 1
            row_n = skip_rows
            while row_n < len(rows):
                t = rows[row_n][col_excel_n].value
                if not t:
                    row_n += 1
                    continue
                if any(s == t for s in v):
                    del rows[row_n]
                else:
                    row_n += 1

        for r in rows:
            ws.append([cell.value for cell in r])


def main():
    config, users = parse_config(CONFIG_PATH)
    src = config["文件路径"]
    src_dir = os.path.dirname(src)
    stem, ext = os.path.splitext(os.path.basename(src))

    # 每个用户单独生成一份：<源文件名>_<用户>.xlsx，放在源文件同目录
    for user_name, user_cfg in users.items():
        out = os.path.join(src_dir, f"{stem}_{user_name}{ext}")
        c = Handle(src)
        c.wb = load_workbook(src)
        for sheet in user_cfg["_sheets"]:
            c.del_column(
                sh_name=sheet["name"],
                skip_rows=sheet["skip_rows"],
                inc=sheet["inc"],
                excl=sheet["excl"],
                excl_eq=sheet["excl_eq"],
            )
        c.wb.save(out)
        c.wb.close()
        print(f"已生成: {out}")

    # 删除视觉空行：选中内容行末尾第一列的单元格，ctrl+shift+方向键下，右键删除，选择整行，保存。留意到右侧滑块变长了


if __name__ == '__main__':
    main()
